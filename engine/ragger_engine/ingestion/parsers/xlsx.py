"""XLSX spreadsheet parser preserving multi-sheet workbooks into DatasetModel."""

from pathlib import Path
from typing import Any, Dict, List
import openpyxl

from ragger_engine.ingestion.exceptions import ParserError
from ragger_engine.ingestion.models import (
    ColumnSchema,
    DatasetModel,
    SourceRecord,
    TableSheet,
)
from ragger_engine.ingestion.parsers.base import BaseParser
from ragger_engine.ingestion.security import ArchiveSecurityValidator


class XLSXParser(BaseParser):
    """Parses Excel XLSX workbooks preserving sheet identity into DatasetModel.

    Uses data_only=True to extract stored/cached calculation values without
    evaluating Excel formula engines locally.
    """

    def parse(self, file_path: Path, source: SourceRecord) -> DatasetModel:
        # Enforce archive decompression constraints
        ArchiveSecurityValidator.validate_zip_archive(file_path)

        try:
            # openpyxl with data_only=True reads cached calculated values if present
            wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        except Exception as e:
            raise ParserError(f"Failed to open XLSX workbook: {e}", parser_name="XLSXParser")

        sheets: List[TableSheet] = []
        warnings: List[str] = []
        total_rows_all = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_rows: List[Dict[str, Any]] = []
            headers: List[str] = []
            columns_data: Dict[str, List[Any]] = {}

            row_iter = ws.iter_rows(values_only=True)
            first_row = next(row_iter, None)

            if not first_row:
                # Empty sheet
                sheets.append(
                    TableSheet(
                        sheet_name=sheet_name,
                        columns=[],
                        rows=[],
                        total_rows=0,
                        total_columns=0,
                        warnings=["Worksheet is empty."],
                    )
                )
                continue

            # Process header row
            headers = [
                str(c).strip() if c is not None and str(c).strip() else f"col_{i}"
                for i, c in enumerate(first_row)
            ]
            columns_data = {h: [] for h in headers}
            sheet_row_count = 0

            for row_values in row_iter:
                sheet_row_count += 1
                row_dict = {}
                # Match row length with headers
                padded = list(row_values) + [None] * (len(headers) - len(row_values))

                for h, val in zip(headers, padded):
                    row_dict[h] = val
                    if sheet_row_count <= 20_000:
                        columns_data[h].append(val)

                if sheet_row_count <= 1000:
                    sheet_rows.append(row_dict)

            total_rows_all += sheet_row_count

            # Build column schema for this sheet
            col_schemas: List[ColumnSchema] = []
            for h in headers:
                vals = columns_data[h]
                non_nulls = [v for v in vals if v is not None]
                null_count = len(vals) - len(non_nulls)
                distinct_count = len(set(str(v) for v in non_nulls))
                sample_vals = non_nulls[:5]

                # Determine dominant type
                dom_type = "string"
                num_vals = []
                for v in non_nulls:
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        num_vals.append(v)

                if len(num_vals) == len(non_nulls) and non_nulls:
                    dom_type = "integer" if all(isinstance(v, int) for v in num_vals) else "float"
                elif any(isinstance(v, bool) for v in non_nulls):
                    dom_type = "boolean"

                min_v = min(num_vals) if num_vals else None
                max_v = max(num_vals) if num_vals else None
                mean_v = sum(num_vals) / len(num_vals) if num_vals else None

                col_schemas.append(
                    ColumnSchema(
                        name=h,
                        inferred_type=dom_type,
                        null_count=null_count,
                        distinct_count=distinct_count,
                        sample_values=sample_vals,
                        min_value=min_v,
                        max_value=max_v,
                        mean_value=round(mean_v, 2) if mean_v is not None else None,
                    )
                )

            sheets.append(
                TableSheet(
                    sheet_name=sheet_name,
                    columns=col_schemas,
                    rows=sheet_rows,
                    total_rows=sheet_row_count,
                    total_columns=len(headers),
                )
            )

        wb.close()

        return DatasetModel(
            source_id=source.source_id,
            title=file_path.stem,
            format="xlsx",
            metadata={"sheet_count": len(sheets)},
            sheets=sheets,
            total_sheets=len(sheets),
            total_rows_across_sheets=total_rows_all,
            warnings=warnings,
        )
