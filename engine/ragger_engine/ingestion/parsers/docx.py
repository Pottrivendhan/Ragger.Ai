"""DOCX format parser preserving heading hierarchy, paragraphs, lists, and tables."""

import uuid
from pathlib import Path
from typing import List
import docx

from ragger_engine.ingestion.exceptions import ParserError
from ragger_engine.ingestion.models import (
    DocumentBlock,
    DocumentModel,
    DocumentTable,
    SourceLocation,
    SourceRecord,
)
from ragger_engine.ingestion.parsers.base import BaseParser
from ragger_engine.ingestion.security import ArchiveSecurityValidator


class DOCXParser(BaseParser):
    """Parses DOCX documents into a structured DocumentModel."""

    def parse(self, file_path: Path, source: SourceRecord) -> DocumentModel:
        # Enforce archive security checks before extraction
        ArchiveSecurityValidator.validate_zip_archive(file_path)

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ParserError(f"Failed to parse DOCX file: {e}", parser_name="DOCXParser")

        blocks: List[DocumentBlock] = []
        tables: List[DocumentTable] = []
        warnings: List[str] = []
        total_words = 0
        block_idx = 0

        # Traverse body elements (paragraphs and tables)
        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                # Find matching python-docx paragraph object
                p_objs = [p for p in doc.paragraphs if p._element == element]
                p_obj = p_objs[0] if p_objs else None
                p_text = p_obj.text.strip() if p_obj else ""
                if not p_text:
                    continue

                words = p_text.split()
                word_count = len(words)
                total_words += word_count

                # Detect style and heading hierarchy
                style_name = ""
                if p_obj and p_obj.style:
                    style_name = p_obj.style.name.lower()

                level = None
                block_type = "paragraph"

                if "heading 1" in style_name:
                    block_type = "heading"
                    level = 1
                elif "heading 2" in style_name:
                    block_type = "heading"
                    level = 2
                elif "heading 3" in style_name:
                    block_type = "heading"
                    level = 3
                elif "heading" in style_name:
                    block_type = "heading"
                    level = 4
                elif "list" in style_name or "bullet" in style_name:
                    block_type = "list_item"

                blocks.append(
                    DocumentBlock(
                        block_id=f"blk_{uuid.uuid4().hex[:8]}",
                        type=block_type,
                        content=p_text,
                        level=level,
                        location=SourceLocation(block_index=block_idx),
                        metadata={"style": style_name, "word_count": word_count},
                    )
                )
                block_idx += 1

            elif tag == "tbl":
                # Find matching python-docx table object
                matching_tables = [t for t in doc.tables if t._element == element]
                if not matching_tables:
                    continue
                tbl = matching_tables[0]

                table_rows: List[List[str]] = []
                for row in tbl.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_rows.append(row_cells)

                if table_rows:
                    headers = table_rows[0]
                    data_rows = table_rows[1:] if len(table_rows) > 1 else []

                    doc_table = DocumentTable(
                        headers=headers,
                        rows=data_rows,
                        location=SourceLocation(block_index=block_idx),
                    )
                    tables.append(doc_table)

                    # Build Markdown representation for searchability
                    md_lines = [" | ".join(headers), " | ".join(["---"] * len(headers))]
                    for r in data_rows:
                        md_lines.append(" | ".join(r))
                    md_content = "\n".join(md_lines)

                    blocks.append(
                        DocumentBlock(
                            block_id=f"blk_{uuid.uuid4().hex[:8]}",
                            type="table",
                            content=md_content,
                            location=SourceLocation(block_index=block_idx),
                            metadata={"rows": len(table_rows), "cols": len(headers)},
                        )
                    )
                    block_idx += 1

        title = file_path.stem
        if doc.core_properties and doc.core_properties.title:
            title = doc.core_properties.title

        return DocumentModel(
            source_id=source.source_id,
            title=str(title),
            format="docx",
            metadata={
                "author": doc.core_properties.author if doc.core_properties else None,
                "created": str(doc.core_properties.created) if doc.core_properties and doc.core_properties.created else None,
            },
            blocks=blocks,
            tables=tables,
            total_blocks=len(blocks),
            total_words=total_words,
            warnings=warnings,
        )
