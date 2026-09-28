"""PPTX presentation parser preserving slides, titles, shapes, notes, and tables."""

import uuid
from pathlib import Path
from typing import List
from pptx import Presentation

from ragger_engine.ingestion.exceptions import ParserError
from ragger_engine.ingestion.models import (
    DocumentBlock,
    DocumentModel,
    DocumentTable,
    SourceLocation,
    SourceRecord,
)
from ragger_engine.ingestion.parsers.base import BaseParser
from ragger_engine.ingestion.security import ArchiveSecurityValidator


class PPTXParser(BaseParser):
    """Parses PPTX presentations into a structured DocumentModel."""

    def parse(self, file_path: Path, source: SourceRecord) -> DocumentModel:
        ArchiveSecurityValidator.validate_zip_archive(file_path)

        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            raise ParserError(f"Failed to parse PPTX file: {e}", parser_name="PPTXParser")

        blocks: List[DocumentBlock] = []
        tables: List[DocumentTable] = []
        warnings: List[str] = []
        total_words = 0
        block_idx = 0

        total_slides = len(prs.slides)
        if total_slides == 0:
            warnings.append("Presentation contains 0 slides.")

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()

            # Record slide header
            if slide_title:
                blocks.append(
                    DocumentBlock(
                        block_id=f"blk_{uuid.uuid4().hex[:8]}",
                        type="heading",
                        content=slide_title,
                        level=2,
                        location=SourceLocation(slide_number=slide_idx, block_index=block_idx),
                        metadata={"slide": slide_idx, "is_slide_title": True},
                    )
                )
                block_idx += 1
                total_words += len(slide_title.split())

            # Traverse slide shapes
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue  # Already captured as slide title

                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        paragraphs = [p.strip() for p in text.splitlines() if p.strip()]
                        for p in paragraphs:
                            words = p.split()
                            word_count = len(words)
                            total_words += word_count

                            blocks.append(
                                DocumentBlock(
                                    block_id=f"blk_{uuid.uuid4().hex[:8]}",
                                    type="paragraph",
                                    content=p,
                                    location=SourceLocation(slide_number=slide_idx, block_index=block_idx),
                                    metadata={"slide": slide_idx, "word_count": word_count},
                                )
                            )
                            block_idx += 1

                elif shape.has_table:
                    tbl = shape.table
                    tbl_rows: List[List[str]] = []
                    for row in tbl.rows:
                        row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        tbl_rows.append(row_text)

                    if tbl_rows:
                        headers = tbl_rows[0]
                        data_rows = tbl_rows[1:] if len(tbl_rows) > 1 else []

                        tables.append(
                            DocumentTable(
                                headers=headers,
                                rows=data_rows,
                                location=SourceLocation(slide_number=slide_idx, block_index=block_idx),
                            )
                        )

                        md_lines = [" | ".join(headers), " | ".join(["---"] * len(headers))]
                        for r in data_rows:
                            md_lines.append(" | ".join(r))

                        blocks.append(
                            DocumentBlock(
                                block_id=f"blk_{uuid.uuid4().hex[:8]}",
                                type="table",
                                content="\n".join(md_lines),
                                location=SourceLocation(slide_number=slide_idx, block_index=block_idx),
                                metadata={"slide": slide_idx, "rows": len(tbl_rows), "cols": len(headers)},
                            )
                        )
                        block_idx += 1

            # Extract Speaker Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    blocks.append(
                        DocumentBlock(
                            block_id=f"blk_{uuid.uuid4().hex[:8]}",
                            type="quote",
                            content=f"Speaker Notes: {notes_text}",
                            location=SourceLocation(slide_number=slide_idx, block_index=block_idx),
                            metadata={"slide": slide_idx, "is_speaker_note": True},
                        )
                    )
                    block_idx += 1
                    total_words += len(notes_text.split())

            # Slide Break marker
            if slide_idx < total_slides:
                blocks.append(
                    DocumentBlock(
                        block_id=f"blk_{uuid.uuid4().hex[:8]}",
                        type="page_break",
                        content=f"--- Slide {slide_idx} End ---",
                        location=SourceLocation(slide_number=slide_idx, block_index=block_idx),
                        metadata={"slide": slide_idx},
                    )
                )
                block_idx += 1

        title = file_path.stem
        if prs.core_properties and prs.core_properties.title:
            title = prs.core_properties.title

        return DocumentModel(
            source_id=source.source_id,
            title=str(title),
            format="pptx",
            metadata={
                "total_slides": total_slides,
                "author": prs.core_properties.author if prs.core_properties else None,
            },
            blocks=blocks,
            tables=tables,
            total_blocks=len(blocks),
            total_words=total_words,
            warnings=warnings,
        )
